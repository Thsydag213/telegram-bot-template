"""
db/exports.py — выгрузка участников и результатов в Excel и CSV.

Колонки заданий есть только при включённом бинго. Время — в часовом
поясе ивента.
"""

import csv
import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from config import FEATURE_BINGO, STATUS_LABELS, TASKS, WIN_THRESHOLD
from db.game import all_completions
from db.users import all_users
from services.timefmt import to_local

_DONE_FILL = PatternFill("solid", fgColor="C6EFCE")
_UNDONE_FILL = PatternFill("solid", fgColor="FFCCCC")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")


def _headers() -> list[str]:
    headers = ["user_id", "username", "имя в Telegram", "телефон",
               "статус сверки", "комментарий сверки", "первый вход", "последний вход"]
    if FEATURE_BINGO:
        for n, name in TASKS.items():
            headers += [f"{n}. {name[:25]}", f"{n}: когда засчитано"]
        headers += ["всего выполнено", "время победы"]
    return headers


async def _rows() -> list[list]:
    users = await all_users()
    completions = await all_completions() if FEATURE_BINGO else {}
    rows = []
    for u in users:
        row = [
            u["user_id"],
            f"@{u['username']}" if u["username"] else "",
            u["full_name"] or "",
            u["phone"] or "",
            STATUS_LABELS.get(u["status"], u["status"] or ""),
            u["status_comment"] or "",
            to_local(u["first_seen"]),
            to_local(u["last_seen"]),
        ]
        if FEATURE_BINGO:
            done = completions.get(u["user_id"], {})
            for n in TASKS:
                row += ["✓" if n in done else "✗", to_local(done.get(n))]
            row += [len(done), to_local(u["ts_win"])]
        rows.append(row)
    # Победители первыми, в порядке времени победы
    if FEATURE_BINGO:
        rows.sort(key=lambda r: (r[-1] == "", r[-1], -r[-2]))
    return rows


async def export_csv() -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_headers())
    writer.writerows(await _rows())
    # BOM — чтобы Excel открыл кириллицу без танцев с кодировкой
    return ("﻿" + buf.getvalue()).encode("utf-8")


async def export_xlsx() -> bytes:
    headers = _headers()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Участники"
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 36

    first_task_col = 9  # после восьми колонок участника
    for row in await _rows():
        ws.append(row)
        if not FEATURE_BINGO:
            continue
        r = ws.max_row
        for i in range(len(TASKS)):
            cell = ws.cell(row=r, column=first_task_col + i * 2)
            cell.fill = _DONE_FILL if cell.value == "✓" else _UNDONE_FILL
            cell.alignment = Alignment(horizontal="center")
        total = ws.cell(row=r, column=first_task_col + len(TASKS) * 2)
        if isinstance(total.value, int) and total.value >= WIN_THRESHOLD:
            total.font = Font(bold=True, color="375623")
        win = ws.cell(row=r, column=first_task_col + len(TASKS) * 2 + 1)
        if win.value:
            win.font = Font(bold=True, color="C00000")

    for col in ws.iter_cols(min_col=1, max_col=ws.max_column, max_row=1):
        header = str(col[0].value or "")
        letter = col[0].column_letter
        if "когда" in header or "время" in header or "вход" in header:
            ws.column_dimensions[letter].width = 20
        elif header[:1].isdigit():
            ws.column_dimensions[letter].width = 12
        else:
            ws.column_dimensions[letter].width = 18
    ws.freeze_panes = "C2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
